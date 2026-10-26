# -*- coding: utf-8 -*-
"""Fix wholesale prices that are higher than retail in products catalog."""

import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from openpyxl import load_workbook
import copy

INPUT_FILE = r"G:\ai\Naggar Print POS Agent\exports\products_catalog_master.xlsx"

# Correct wholesale prices (must be LOWER than selling_price/retail)
CORRECTIONS = {
    # id: (selling_price/retail, current_wholesale, new_wholesale, reason)
    51:    (330, 380, 260, "composite box: sum of components - 5% bundle discount"),
    10009: (90,  100,  65, "graduation cap: ~28% below retail"),
    10011: (150, 200, 115, "royal shawl: ~23% below retail"),
    10034: (150, 170, 120, "black shawl: ~20% below retail"),
}

print("Loading workbook...")
wb = load_workbook(INPUT_FILE)
ws = wb.active

print(f"Sheet name: {ws.title}")
print(f"Total rows: {ws.max_row}")

# Find the column indices (id is col A, notes is col N)
id_col = None
notes_col = None
name_col = None
selling_price_col = None

for col in range(1, ws.max_column + 1):
    val = ws.cell(row=2, column=col).value
    if val == "id":
        id_col = col
    elif val == "name":
        name_col = col
    elif val == "selling_price":
        selling_price_col = col
    elif val == "notes":
        notes_col = col

print(f"Columns: id={id_col}, name={name_col}, selling_price={selling_price_col}, notes={notes_col}")

# Find and update each product
updated = []
not_found = []

for row in range(4, ws.max_row + 1):
    cell_id = ws.cell(row=row, column=id_col).value
    if cell_id is None:
        continue
    
    # Handle both string and int IDs
    try:
        product_id = int(cell_id)
    except (ValueError, TypeError):
        continue
    
    if product_id in CORRECTIONS:
        retail, current_wholesale, new_wholesale, reason = CORRECTIONS[product_id]
        name = ws.cell(row=row, column=name_col).value or "Unknown"
        
        # Get current notes
        current_notes = ws.cell(row=row, column=notes_col).value or ""
        
        # Build new wholesale note
        new_note = f"جملة: {new_wholesale}.0 ج (approved: Fares — {reason})"
        
        # Update the notes column
        ws.cell(row=row, column=notes_col).value = new_note
        
        print(f"\n✓ id={product_id} | {name}")
        print(f"  Retail: {retail} | Old wholesale: {current_wholesale} | New wholesale: {new_wholesale}")
        print(f"  Reason: {reason}")
        print(f"  Old notes: {current_notes}")
        print(f"  New notes: {new_note}")
        
        updated.append(product_id)
        
        # Verify the new wholesale is lower than retail
        if new_wholesale >= retail:
            print(f"  ⚠️ WARNING: New wholesale ({new_wholesale}) is NOT lower than retail ({retail})!")
        else:
            print(f"  ✓ Verified: {new_wholesale} < {retail}")

print(f"\n\nSummary:")
print(f"  Updated: {len(updated)} products")
print(f"  Not found: {not_found}")

# Save the file
print(f"\nSaving to {INPUT_FILE}...")
wb.save(INPUT_FILE)
print("Done!")
