# -*- coding: utf-8 -*-
import sys
try:
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
except Exception:
    pass
from openpyxl import load_workbook

wb = load_workbook(r'C:\Users\mahmo\Downloads\كتالوج_النجار.xlsx')
print('sheets:', wb.sheetnames)
ws = wb.active
print('active:', ws.title, '| dims:', ws.dimensions, '| max_row:', ws.max_row, '| max_col:', ws.max_column)
print()
for r in range(1, 7):
    vals = [ws.cell(row=r, column=j).value for j in range(1, min(ws.max_column, 15) + 1)]
    print(f'row {r}:', vals)
print()
print('--- rows 106..112 (tail) ---')
for r in range(106, min(ws.max_row, 112) + 1):
    vals = [ws.cell(row=r, column=j).value for j in range(1, min(ws.max_column, 15) + 1)]
    print(f'row {r}:', vals)
